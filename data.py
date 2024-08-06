import gdown
import os
import time
import pandas as pd
from openpyxl import load_workbook
from typing import Dict, Optional

class GameManager:
    def __init__(self, file_id: Optional[str] = None, output_path: str = "data.xlsx"):
        self.file_id = file_id
        self.output_path = output_path
        self.dataframes: Dict[str, pd.DataFrame] = {}

    def download_data(self) -> bool:
        """Download the Excel file from Google Drive."""
        if not self.file_id:
            print("No file_id provided. Cannot download data.")
            return False
        
        url = f"https://drive.google.com/uc?id={self.file_id}"
        
        if os.path.exists(self.output_path):
            print(f"File already exists at {self.output_path}")
            return True

        gdown.download(url, self.output_path, quiet=False)
        
        return self._wait_for_download_completion()

    def _wait_for_download_completion(self) -> bool:
        """Wait for the file to be completely downloaded."""
        while not os.path.exists(self.output_path):
            time.sleep(1)

        file_size = 0
        while True:
            current_size = os.path.getsize(self.output_path)
            if current_size == file_size:
                print("Download completed.")
                return True
            file_size = current_size
            time.sleep(1)

    def get_table_data(self, target_table: str = 'RESULTS_LOG') -> pd.DataFrame:
        """Read and return the target table from the Excel file."""
        workbook = load_workbook(self.output_path, data_only=True)
        sheet = workbook.active

        for table_name, table_range in sheet.tables.items():
            if table_name == target_table:
                print(f"Reading Table: {table_name}")
                data = [[cell.value for cell in row] for row in sheet[table_range]]
                return pd.DataFrame(data[1:], columns=data[0])

        print(f"Table '{target_table}' not found.")
        return pd.DataFrame()  # Return an empty DataFrame if table not found
    

